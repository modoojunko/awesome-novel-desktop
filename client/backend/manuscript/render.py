"""成稿三格式渲染器（c-manuscript-download PR2）。

只消费 content.py 的 Manuscript 结构，零查库零 IO（bytes/str 进出）。
- md：`# 书名` / `## 第X卷` / `### 第X章` 层级，段落空行分隔；
- txt：纯文本，卷/章标题行 + 空行段落；
- docx：python-docx heading/paragraph（依赖为既有 requirements 项，
  惰性 import——缺失时报 JobError 而非崩进程，novels/importer.py 同先例）。

序号标签口径与前端 nodeTitle 同族：默认序号形态（「第三章」）不重复拼前缀。
"""

import re

from job_runner import JobError
from manuscript.content import Manuscript

_CN = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九"]


def cn_num(n: int) -> str:
    """与前端 nodeTitle.ts cnNum 同族（1 → 一，12 → 十二，21 → 二十一，102 → 一百二）。

    四位及以上回退阿拉伯数字（1000 → "1000"）：中文数字只排到百位，前端同族实现同为
    回退。修复前这里直接 `_CN[hundreds]` 越界——单卷满千章时 IndexError 会让**整单
    下载**（md/txt/docx 全格式）硬失败（P3，2026-09-18）。
    """
    if n <= 0 or not isinstance(n, int):
        return str(n)
    if n >= 1000:
        return str(n)
    if n < 10:
        return _CN[n]
    if n < 20:
        return "十" + (_CN[n % 10] if n % 10 else "")
    if n < 100:
        return _CN[n // 10] + "十" + (_CN[n % 10] if n % 10 else "")
    hundreds = n // 100
    rest = n % 100
    return _CN[hundreds] + "百" + (cn_num(rest) if rest else "")


_DEFAULT_TITLE_RE = re.compile(r"^第\s*[0-9一二三四五六七八九十百零]+\s*[卷章]$")  # \s* 与前端 nodeTitle 容差一致


def _label(kind: str, no: int, title: str) -> str:
    """「第X章 · 名」；title 已是任意默认序号形态（「第三章」「第3章」）时不重复
    拼（nodeTitle 同口径——「第二章 · 第三章」是用户没起名的常见态）。"""
    prefix = f"第{cn_num(no)}{kind}"
    name = (title or "").strip()
    if not name or name == prefix or _DEFAULT_TITLE_RE.match(name):
        return prefix
    return f"{prefix} · {name}"


def _prose_lines(prose: str) -> list[str]:
    return [ln.strip() for ln in (prose or "").splitlines() if ln.strip()]


def render_md(m: Manuscript) -> str:
    out: list[str] = [f"# {m.title}", ""]
    for vol in m.volumes:
        out += [f"## {_label('卷', vol.no, vol.title)}", ""]
        for ch in vol.chapters:
            out += [f"### {_label('章', ch.no, ch.title)}", ""]
            for ln in _prose_lines(ch.prose):
                out += [ln, ""]
    return "\n".join(out).rstrip() + "\n"


def render_txt(m: Manuscript) -> str:
    out: list[str] = [m.title, ""]
    for vol in m.volumes:
        out += [f"◇ {_label('卷', vol.no, vol.title)}", ""]
        for ch in vol.chapters:
            out += [_label("章", ch.no, ch.title), ""]
            for ln in _prose_lines(ch.prose):
                out += [f"　　{ln}", ""]
    return "\n".join(out).rstrip() + "\n"


def render_docx(m: Manuscript) -> bytes:
    try:
        import io

        import docx
    except ImportError as e:  # 打包漏收集时的可读失败（不崩进程）
        raise JobError("docx_unavailable", "Word 渲染组件不可用，请重新安装应用") from e

    document = docx.Document()
    document.add_heading(m.title, level=0)
    for vol in m.volumes:
        document.add_heading(_label("卷", vol.no, vol.title), level=1)
        for ch in vol.chapters:
            document.add_heading(_label("章", ch.no, ch.title), level=2)
            for ln in _prose_lines(ch.prose):
                document.add_paragraph(ln)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
